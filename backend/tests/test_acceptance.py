import os, tempfile, json, time
os.environ['EDUPULSE_DATA_DIR']=tempfile.mkdtemp(prefix='edupulse-tests-')
import pytest
from fastapi.testclient import TestClient
from datetime import datetime, timezone, timedelta
from io import BytesIO
from openpyxl import load_workbook
from PIL import Image
from app.main import app, failures
from app.db import *
from app.security import ph

@pytest.fixture(autouse=True)
def reset():
    Base.metadata.drop_all(engine); init_db(); failures.clear()
    with SessionLocal.begin() as db:
        db.add(User(username='test-admin',display_name='测试主管',role='admin',password_hash=ph.hash('TestOnlyPassword2026')))

def client(username='test-admin',password='TestOnlyPassword2026'):
    c=TestClient(app)
    r=c.post('/api/auth/login',json={'username':username,'password':password})
    assert r.status_code==200,r.text
    c.headers['x-csrf-token']=r.json()['csrf_token']; return c

def setup(n=2):
    a=client(); d=a.get('/api/dictionaries').json(); campus=d['campus'][0]['id']; math=d['subject'][0]['id']; english=d['subject'][1]['id']
    teachers=[]
    for i in range(n):
        u=a.post('/api/admin/teachers',json={'display_name':'同名老师' if i<2 else '测试老师','campus_id':campus}).json()
        p=a.post(f"/api/admin/teachers/{u['id']}/issue-password").json()
        teachers.append((u,client(p['username'],p['password']),p))
    return a,teachers,campus,math,english

def student(t,campus,name='虚构学生'):
    r=t.post('/api/students',json={'display_name':name,'grade':'高一','campus_id':campus}); assert r.status_code==200,r.text; return r.json()['id']
def classroom(t,subject,capacity,ids,slots=1):
    body={'display_name':'测试周三班','subject_id':subject,'grade':'高一','capacity':capacity,'students':[{'student_id':s,'level':'A','note':'仅测试评价','weekly_slots':[{'weekday':i+1,'slot_number':3} for i in range(slots)]} for s in ids]}
    r=t.post('/api/teacher/classes',json=body); assert r.status_code==200,r.text; return r.json()
def broadcast(a,campus,ids=None):
    r=a.post('/api/admin/broadcasts',json={'title':'测试核对','campus_id':campus,'teacher_ids':ids or [],'starts_at':(datetime.now(timezone.utc)-timedelta(minutes=1)).isoformat(),'deadline_at':(datetime.now(timezone.utc)+timedelta(hours=1)).isoformat()}); assert r.status_code==200,r.text; return r.json()['id']
def body_of(c):
    return {k:c[k] for k in ('display_name','subject_id','grade','capacity','version')}|{'students':[{k:s[k] for k in ('student_id','level','note')}|{'weekly_slots':[{k:slot[k] for k in ('weekday','slot_number','start_time','end_time')} for slot in s['weekly_slots']]} for s in c['students']]}

def test_accounts_isolation_reset_disable_and_no_plaintext():
    a,ts,campus,math,eng=setup(); u,t,p=ts[0]; v,other,p2=ts[1]
    assert p['password']!=p2['password']; assert len(p['password'])>=14; assert u['username']!=v['username']
    c=classroom(other,math,3,[student(other,campus)])
    assert t.get('/api/teacher/classes').json()==[]
    for method,path,kwargs in [('get',f"/api/teacher/classes/{c['id']}",{}),('patch',f"/api/teacher/classes/{c['id']}",{'json':body_of(c)}),('delete',f"/api/teacher/classes/{c['id']}?confirmed=true&version=1",{})]:
        assert getattr(t,method)(path,**kwargs).status_code==404
    assert t.get('/api/admin/teachers').status_code==403
    assert t.post(f"/api/admin/teachers/{v['id']}/issue-password").status_code==403
    assert t.get('/api/admin/reports/export.xlsx').status_code==403
    assert TestClient(app).get('/api/teacher/classes').status_code==401
    payload=body_of(c)|{'teacher_id':v['id']}; assert t.post('/api/teacher/classes',json=payload).status_code==422
    before=a.get('/api/admin/teachers').json(); assert all('password' not in x and 'password_hash' not in x for x in before)
    new=a.post(f"/api/admin/teachers/{u['id']}/issue-password").json()
    assert t.get('/api/teacher/classes').status_code==401
    assert TestClient(app).post('/api/auth/login',json={'username':p['username'],'password':p['password']}).status_code==401
    assert other.get('/api/teacher/classes').status_code==200
    t=client(new['username'],new['password'])
    a.patch(f"/api/admin/teachers/{u['id']}",json={'display_name':u['display_name'],'campus_id':campus,'enabled':False})
    assert t.get('/api/teacher/classes').status_code==401
    with SessionLocal() as db:
        assert db.get(User,u['id']).password_hash.startswith('$argon2id$')
        for table in (Audit,Snapshot):
            values=str([vars(x) for x in db.query(table)])
            assert p['password'] not in values and new['password'] not in values

def test_acceptance_weighted_occupancy_workload_and_student_dedup():
    a,ts,campus,math,eng=setup(); u,t,_=ts[0]; other=ts[1][1]
    shared=student(t,campus,'同名学生'); distinct=student(t,campus,'同名学生')
    for cap,count in [(3,1),(3,2),(3,3),(6,2),(6,4)]:
        classroom(t,math,cap,[shared]+[student(t,campus) for _ in range(count-1)],slots=2 if cap==6 and count==4 else 1)
    classroom(t,math,1,[distinct]); classroom(other,eng,3,[shared])
    data=a.get('/api/admin/reports').json(); r=next(x for x in data['teachers'] if x['id']==u['id'])
    assert r['groups']['3']['rate']==pytest.approx(6/9)
    assert r['groups']['6']['rate']==pytest.approx(6/12)
    assert r['small_rate']==pytest.approx(12/21)
    assert r['groups']['6']['weekly_lessons']==3
    assert r['weekly_lessons']==7
    r2=next(x for x in data['teachers'] if x['id']!=u['id']); assert r2['groups']['6']['rate'] is None
    s=next(x for x in data['students'] if x['id']==shared); assert s['subject_count']==2; assert len(s['details'])==6
    assert len([x for x in data['students'] if x['display_name']=='同名学生'])==2
    results=t.get('/api/students/search?q=同名').json(); assert set(results[0])=={'id','display_name','grade'}
    scoped=a.get(f'/api/admin/reports?subject_id={eng}').json(); assert len(scoped['teachers'])==1

def test_six_capacity_four_students_two_slots_is_two_lessons():
    a,ts,campus,math,_=setup(); t=ts[0][1]
    c=classroom(t,math,6,[student(t,campus) for _ in range(4)],slots=2)
    classroom(t,math,1,[student(t,campus)])
    b=body_of(c); b['display_name']='名称任意修改'; b['students'][0]['note']='更新评价'; b['students'][0]['level']='C'
    r=t.patch(f"/api/teacher/classes/{c['id']}",json=b); assert r.status_code==200,r.text
    fresh=t.get('/api/teacher/classes').json()[0]; assert fresh['students'][0]['note']=='更新评价'; assert fresh['students'][0]['level']=='C'
    assert t.patch(f"/api/teacher/classes/{c['id']}",json=b).status_code==409
    data=a.get('/api/admin/reports/workload').json()[0]; assert data['weekly_lessons']==3; assert data['groups']['6']['weekly_lessons']==2
    b=body_of(fresh); b['capacity']=3; assert t.patch(f"/api/teacher/classes/{c['id']}",json=b).status_code==422
    b=body_of(fresh); b['students'][1]['student_id']=b['students'][0]['student_id']; assert t.patch(f"/api/teacher/classes/{c['id']}",json=b).status_code==422

def test_broadcast_progress_frozen_targets_offline_snapshot_exports():
    a,ts,campus,math,_=setup(3); t=ts[0][1]; c=classroom(t,math,3,[student(t,campus)])
    b=broadcast(a,campus)
    a.post('/api/admin/teachers',json={'display_name':'后来老师','campus_id':campus})
    assert len(a.get(f'/api/admin/broadcasts/{b}/progress').json()['targets'])==3
    # Save does not confirm.
    body=body_of(c); body['display_name']='调整名称'; c=t.patch(f"/api/teacher/classes/{c['id']}",json=body).json()
    assert not any(x['confirmed_at'] for x in a.get(f'/api/admin/broadcasts/{b}/progress').json()['targets'])
    for _,tc,_ in ts[:2]: assert tc.post(f'/api/broadcasts/{b}/confirm').status_code==200
    targets=a.get(f'/api/admin/broadcasts/{b}/progress').json()['targets']; assert sum(bool(x['confirmed_at']) for x in targets)==2
    # Offline login still reads current task.
    p=ts[2][2]; offline=client(p['username'],p['password']); assert offline.get('/api/broadcasts/current').json()[0]['id']==b
    assert t.post(f'/api/broadcasts/{b}/confirm').status_code==200
    assert offline.post(f'/api/broadcasts/{b}/confirm').status_code==200
    snapshot=a.get(f'/api/admin/broadcasts/{b}/snapshot').json()
    body=body_of(c); body['students'].append({'student_id':student(t,campus),'level':'B','note':'后续调整','weekly_slots':[{'weekday':1,'slot_number':3}]})
    assert t.patch(f"/api/teacher/classes/{c['id']}",json=body).status_code==200
    assert a.get(f'/api/admin/broadcasts/{b}/snapshot').json()==snapshot
    assert a.get('/api/admin/reports').json()['teachers'][0]['groups']['3']['seats']==2
    assert snapshot['teachers'][0]['groups']['3']['seats']==1
    b2=broadcast(a,campus,[ts[0][0]['id']]); assert t.post(f'/api/broadcasts/{b2}/confirm').status_code==200
    assert a.get(f'/api/admin/broadcasts/{b}/snapshot').json()==snapshot
    output=a.get(f'/api/admin/reports/export.xlsx?broadcast_id={b}'); assert output.status_code==200
    wb=load_workbook(BytesIO(output.content)); assert {'教师满班率','教师周负荷','学生科目教师评价明细'}<=set(wb.sheetnames)
    assert wb['教师满班率']['H2'].number_format=='0.00%'; assert isinstance(wb['教师满班率']['H2'].value,float)
    output=a.get(f'/api/admin/reports/export.png?broadcast_id={b}'); assert output.status_code==200
    assert Image.open(BytesIO(output.content)).size[0]==1500
    engine.dispose(); assert a.get(f'/api/admin/broadcasts/{b}/snapshot').json()==snapshot

def test_deadline_and_unconfirmed_snapshot():
    a,ts,campus,_,_=setup(); b=broadcast(a,campus)
    with SessionLocal.begin() as db: db.get(Broadcast,b).deadline_at=(datetime.now(timezone.utc)-timedelta(seconds=1)).isoformat()
    assert ts[0][1].post(f'/api/broadcasts/{b}/confirm').status_code==409
    data=a.get(f'/api/admin/broadcasts/{b}/snapshot').json()
    assert data['teachers'][0]['confirmation']=='未经本轮确认'
    assert ts[0][1].get('/api/broadcasts/current').json()==[]

def test_csrf_rate_limit_and_websocket_auth_and_realtime():
    a,ts,campus,_,_=setup(); t=ts[0][1]
    assert a.post('/api/admin/teachers',json={'display_name':'X','campus_id':campus},headers={'x-csrf-token':''}).status_code==403
    assert a.post('/api/admin/teachers',json={'display_name':'X','campus_id':campus},headers={'origin':'http://evil.example'}).status_code==403
    with pytest.raises(Exception):
        with TestClient(app).websocket_connect('/ws/events',headers={'origin':'http://testserver'}): pass
    with t.websocket_connect('/ws/events',headers={'origin':'http://testserver'}) as ws:
        b=broadcast(a,campus); assert ws.receive_json()['event']=='broadcast.created'
    assert a.get('/api/admin/teachers').headers['cache-control']=='no-store'
    bad=TestClient(app)
    for _ in range(10): assert bad.post('/api/auth/login',json={'username':'unknown','password':'bad'}).status_code==401
    assert bad.post('/api/auth/login',json={'username':'unknown','password':'bad'}).status_code==429

def test_admin_manual_close_and_merge():
    a,ts,campus,math,eng=setup(); t=ts[0][1]
    s1=student(t,campus); s2=student(t,campus)
    classroom(t,math,3,[s1]); classroom(t,eng,3,[s2])
    assert a.post('/api/admin/students/merge',json={'source_id':s2,'target_id':s1}).status_code==200
    assert a.get('/api/admin/reports/students').json()[0]['subject_count']==2
    b=broadcast(a,campus); assert a.post(f'/api/admin/broadcasts/{b}/close').status_code==200
    assert a.get(f'/api/admin/broadcasts/{b}/snapshot').status_code==200
    assert t.post(f'/api/broadcasts/{b}/confirm').status_code==409

def test_expired_snapshot_is_frozen_before_next_write():
    a,ts,campus,math,_=setup();t=ts[0][1]
    c=classroom(t,math,3,[student(t,campus)])
    b=broadcast(a,campus)
    # Prepare a student before the deadline, then edit only after expiry.
    sid=student(t,campus)
    with SessionLocal.begin() as db: db.get(Broadcast,b).deadline_at=(datetime.now(timezone.utc)-timedelta(seconds=1)).isoformat()
    body=body_of(c);body['students'].append({'student_id':sid,'level':'B','note':'','weekly_slots':[{'weekday':1,'slot_number':3}]})
    assert t.patch(f"/api/teacher/classes/{c['id']}",json=body).status_code==200
    old=a.get(f'/api/admin/broadcasts/{b}/snapshot').json()
    assert old['teachers'][0]['groups']['3']['seats']==1
    assert a.get('/api/admin/reports').json()['teachers'][0]['groups']['3']['seats']==2

def test_deletion_audit_and_historical_evaluations_survive():
    a,ts,campus,math,_=setup();t=ts[0][1]
    c=classroom(t,math,3,[student(t,campus)])
    b=broadcast(a,campus);a.post(f'/api/admin/broadcasts/{b}/close')
    assert t.delete(f"/api/teacher/classes/{c['id']}?version=1").status_code==422
    assert t.delete(f"/api/teacher/classes/{c['id']}?confirmed=true&version=1").status_code==200
    assert t.get('/api/teacher/classes').json()==[]
    assert a.get(f'/api/admin/broadcasts/{b}/snapshot').json()['students'][0]['details'][0]['note']=='仅测试评价'
    assert any(l['action']=='class.delete' for l in a.get('/api/admin/audit').json())

def test_backup_api_integrity_restore_and_excel_formula_escape(tmp_path):
    import subprocess,sys,sqlite3
    from app.exports import xlsx
    a,ts,campus,math,_=setup();t=ts[0][1]
    classroom(t,math,3,[student(t,campus,'=测试公式')])
    output=tmp_path/'backup.sqlite3'
    subprocess.run([sys.executable,'-m','app.cli','backup','--output',str(output)],check=True,capture_output=True)
    with sqlite3.connect(output) as db:
        assert db.execute('PRAGMA integrity_check').fetchone()[0]=='ok'
        assert db.execute('SELECT count(*) FROM classes').fetchone()[0]==1
    with SessionLocal.begin() as db: db.get(Student,1).display_name='后续名称'
    # Isolated test database only; there is no running production server.
    subprocess.run([sys.executable,'-m','app.cli','restore',str(output),'--confirm'],check=True,capture_output=True)
    with SessionLocal() as db: assert db.get(Student,1).display_name=='=测试公式'
    assert a.get('/api/auth/me').status_code==401
    a=client()
    data=a.get('/api/admin/reports').json();wb=load_workbook(BytesIO(xlsx(data)))
    assert wb['学生科目教师评价明细']['B2'].data_type=='s'
    assert wb['学生科目教师评价明细']['B2'].value.startswith("'=测试")

def test_teacher_search_retains_associated_student_counts():
    a,ts,campus,math,_=setup();t=ts[0][1]
    classroom(t,math,3,[student(t,campus,'教学测试生')])
    data=a.get('/api/admin/reports?q=同名老师&search_type=teacher').json()
    assert len(data['teachers'])==2 and len(data['students'])==1
    data=a.get('/api/admin/reports/students?q=教学测试生').json()
    assert len(data)==1 and data[0]['subject_count']==1
