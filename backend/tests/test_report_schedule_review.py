"""Live broadcast selection, unchanged filter contracts and advisory conflicts."""
import json
from io import BytesIO
import pytest
from openpyxl import load_workbook
from test_acceptance import reset, setup, student, classroom, body_of, broadcast, client
from app.db import SessionLocal, Broadcast, Snapshot


@pytest.mark.parametrize('ending', ['manual', 'confirmed', 'expired'])
def test_live_reports_do_not_inherit_finished_broadcast(ending):
    admin, teachers, campus, subject, _ = setup()
    teacher = teachers[0][1]
    room = classroom(teacher, subject, 3, [student(teacher, campus)])
    bid = broadcast(admin, campus)
    if ending == 'manual':
        assert admin.post(f'/api/admin/broadcasts/{bid}/close').status_code == 200
    elif ending == 'confirmed':
        for _, t, _ in teachers:
            assert t.post(f'/api/broadcasts/{bid}/confirm').status_code == 200
    else:
        with SessionLocal.begin() as db:
            db.get(Broadcast, bid).deadline_at = '2000-01-01T00:00:00+00:00'
        assert admin.get(f'/api/admin/broadcasts/{bid}/snapshot').status_code == 200
    with SessionLocal() as db:
        raw = db.query(Snapshot).filter_by(broadcast_id=bid).one().data
    body = body_of(room); body['display_name'] = '结束广播后的新修改'
    assert teacher.patch(f"/api/teacher/classes/{room['id']}", json=body).status_code == 200
    live = admin.get('/api/admin/reports').json()
    assert 'broadcast' not in live
    assert all('confirmation' not in t for t in live['teachers'])
    assert live['classes'][0]['display_name'] == body['display_name']
    assert admin.get(f'/api/admin/broadcasts/{bid}/snapshot').json() == json.loads(raw)
    wb = load_workbook(BytesIO(admin.get('/api/admin/reports/export.xlsx').content))
    assert wb['填报完成清单'].max_row == 1
    assert wb['教师满班率']['P2'].value == '当前数据'
    with SessionLocal() as db:
        assert db.query(Snapshot).filter_by(broadcast_id=bid).one().data == raw


def test_live_reports_find_active_round_behind_newer_closed_round():
    admin, _, campus, _, _ = setup()
    active = broadcast(admin, campus)
    closed = broadcast(admin, campus)
    admin.post(f'/api/admin/broadcasts/{closed}/close')
    assert admin.get('/api/admin/reports').json()['broadcast']['id'] == active


def test_live_reports_exclude_expired_unfrozen_and_future_rounds():
    admin, _, campus, _, _ = setup()
    expired = broadcast(admin, campus)
    future = broadcast(admin, campus)
    with SessionLocal.begin() as db:
        db.get(Broadcast, expired).deadline_at = '2000-01-01T00:00:00+00:00'
        b = db.get(Broadcast, future)
        b.starts_at = '2099-01-01T00:00:00+00:00'; b.deadline_at = '2099-01-02T00:00:00+00:00'
    assert 'broadcast' not in admin.get('/api/admin/reports').json()


def test_existing_grade_student_and_teacher_filter_contracts_stay_unchanged():
    admin, teachers, campus, subject, english = setup()
    t1, t2 = teachers[0][1], teachers[1][1]
    a, b = student(t1, campus, '甲'), student(t1, campus, '乙')
    room = classroom(t1, subject, 3, [a, b]); classroom(t2, english, 3, [a])
    d = admin.get('/api/dictionaries').json()
    other_grade = next(g['name'] for g in d['grade'] if g['name'] != '高一')
    from app.db import Student
    with SessionLocal.begin() as db:
        db.get(Student, b).grade = other_grade
    full = admin.get('/api/admin/reports').json()
    for params in ({'grade': other_grade, 'search_type': 'student'}, {'q': '乙', 'search_type': 'student'}):
        filtered = admin.get('/api/admin/reports', params=params).json()
        assert [s['id'] for s in filtered['students']] == [b]
        assert filtered['teachers'] == full['teachers']
        assert filtered['classes'] == full['classes']
    filtered = admin.get('/api/admin/reports', params={'teacher_id': teachers[0][0]['id'], 'search_type': 'student'}).json()
    assert len(filtered['teachers']) == 1
    assert [c['id'] for c in filtered['classes']] == [room['id']]
    assert all(d['teacher_id'] == teachers[0][0]['id'] for s in filtered['students'] for d in s['details'])
    assert next(s for s in filtered['students'] if s['id'] == a)['subject_count'] == 1
    wb = load_workbook(BytesIO(admin.get('/api/admin/reports/export.xlsx', params={'grade': other_grade, 'search_type': 'student'}).content))
    assert wb['教师周课量'].max_row == 3
    assert wb['学生科目教师评价明细'].max_row == 2


def schedule(t, room, weekday, number, start=None, end=None):
    body = body_of(room)
    for member in body['students']:
        member['weekly_slots'] = [{'weekday': weekday, 'slot_number': number, 'start_time': start, 'end_time': end}]
    response = t.patch(f"/api/teacher/classes/{room['id']}", json=body)
    assert response.status_code == 200, response.text
    return response.json()


@pytest.mark.parametrize('period2,start2,end2,kind', [
    (2, None, None, 'possible_overlap'),
    (5, '09:30', '10:30', 'time_overlap'),
    (2, '10:00', '11:00', None),
    (2, '11:00', '12:00', None),
    (5, None, None, None),
])
def test_conflict_advisories_visible_to_teacher_and_admin_without_blocking_save(period2, start2, end2, kind):
    admin, teachers, campus, subject, _ = setup()
    user, t, _ = teachers[0]
    a = classroom(t, subject, 3, [student(t, campus), student(t, campus)])
    b = classroom(t, subject, 3, [student(t, campus)])
    a = schedule(t, a, 3, 2, '09:00', '10:00')
    b = schedule(t, b, 3, period2, start2, end2)
    other = teachers[1][1]
    other_room = classroom(other, subject, 3, [student(other, campus)])
    schedule(other, other_room, 3, 2, '09:00', '10:00')
    before = t.get('/api/teacher/classes').json()
    own = t.get('/api/teacher/schedule-warnings')
    assert own.status_code == 200, own.text
    warnings = own.json()
    assert len(warnings) == (1 if kind else 0)
    if kind:
        w = warnings[0]
        assert w['kind'] == kind and w['teacher_id'] == user['id']
        assert {c['id'] for c in w['classes']} == {a['id'], b['id']}
        assert 'student_id' not in json.dumps(w)
    assert admin.get('/api/admin/schedule-warnings').json() == warnings
    assert teachers[1][1].get('/api/teacher/schedule-warnings').json() == []
    assert t.get('/api/admin/schedule-warnings').status_code == 403
    assert client().get('/api/teacher/schedule-warnings').status_code == 403
    assert admin.get('/api/admin/reports/workload').json()[0]['weekly_lessons'] == 2
    assert 'schedule_warnings' not in admin.get('/api/admin/reports').json()
    assert t.get('/api/teacher/classes').json() == before
    # Resolving or removing a class clears warnings, and the other teacher is never compared.
    schedule(t, b, 4, period2, start2, end2)
    assert t.get('/api/teacher/schedule-warnings').json() == []


def test_warning_reads_do_not_change_snapshot_format_or_content():
    admin, teachers, campus, subject, _ = setup()
    t = teachers[0][1]
    for _ in range(2):
        classroom(t, subject, 3, [student(t, campus)])
    bid = broadcast(admin, campus); admin.post(f'/api/admin/broadcasts/{bid}/close')
    with SessionLocal() as db:
        raw = db.query(Snapshot).filter_by(broadcast_id=bid).one().data
    assert admin.get('/api/admin/schedule-warnings').status_code == 200
    assert t.get('/api/teacher/schedule-warnings').status_code == 200
    assert admin.get(f'/api/admin/broadcasts/{bid}/snapshot').json() == json.loads(raw)
    with SessionLocal() as db:
        assert db.query(Snapshot).filter_by(broadcast_id=bid).one().data == raw
