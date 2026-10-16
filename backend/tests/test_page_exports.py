from io import BytesIO
import pytest
from openpyxl import load_workbook
from PIL import Image
from test_acceptance import reset,setup,student,classroom,broadcast,body_of


def workbook(admin,query):
    response=admin.get('/api/admin/reports/export.xlsx?'+query)
    assert response.status_code==200,response.text
    return load_workbook(BytesIO(response.content))


def test_exports_match_page_filters_order_and_current_metrics():
    a,ts,campus,math,english=setup();t=ts[0][1]
    first=student(t,campus,'=保留公式安全');second=student(t,campus,'乙');empty=student(t,campus,'未关联学生')
    classroom(t,math,3,[first,second]);classroom(t,english,6,[first])
    wb=workbook(a,'view=occupancy')
    assert set(wb.sheetnames)=={'教师满班率','满班率汇总','统计说明'}
    ws=wb['教师满班率'];row=list(ws.values)[1];assert row[8]==2/3 and row[9]==2
    assert row[14]==pytest.approx(1/6) and row[15]==1 and row[17]==pytest.approx(1/3) and row[18]==1.5
    assert ws['I2'].number_format=='0.00%' and ws['J2'].number_format=='0.00'
    wb=workbook(a,f'view=occupancy&subject_id={math}')
    assert wb['教师满班率']['O2'].value=='—'
    wb=workbook(a,f'view=students&ordered_ids={empty},{second},{first}')
    assert set(wb.sheetnames)=={'学生列表','学生评价与档位明细','统计说明'}
    rows=list(wb['学生列表'].values);assert [r[0] for r in rows[1:]]==[empty,second,first]
    assert rows[1][3:]==(0,None,0)
    assert rows[3][1]=="'=保留公式安全"
    assert wb['学生列表']['B4'].data_type=='s'
    assert wb['学生评价与档位明细'].max_row==4
    wb=workbook(a,f'view=students&subject_id={math}&q=乙')
    assert wb['学生列表'].max_row==2 and wb['学生列表']['A2'].value==second
    assert wb['学生评价与档位明细'].max_row==2
    # Different class types all count; 18 lessons must export as 120%, unbounded.
    c=t.get('/api/teacher/classes').json()[0];body=body_of(c)
    for member in body['students']:member['weekly_slots']=[{'weekday':i//6+1,'slot_number':i%6+1} for i in range(17)]
    assert t.patch(f"/api/teacher/classes/{c['id']}",json=body).status_code==200
    wb=workbook(a,'view=workload');assert set(wb.sheetnames)=={'教师负荷','统计说明'}
    ws=wb['教师负荷'];assert ws['G2'].value==18 and ws['I2'].value=='18/15' and ws['J2'].value==1.2
    assert ws['J2'].number_format=='0.00%'
    for view in ('occupancy','workload','students'):
        r=a.get('/api/admin/reports/export.png?view='+view);assert r.status_code==200
        image=Image.open(BytesIO(r.content));assert image.width>1000 and image.height>300
    assert a.get('/api/admin/reports/export.xlsx?view=invalid').status_code==422
    assert a.get('/api/admin/reports/export.xlsx?view=students&ordered_ids=x').status_code==422
    assert t.get('/api/admin/reports/export.xlsx?view=students').status_code==403


def test_page_exports_preserve_historical_source():
    a,ts,campus,math,_=setup();t=ts[0][1]
    c=classroom(t,math,3,[student(t,campus)])
    bid=broadcast(a,campus);a.post(f'/api/admin/broadcasts/{bid}/close')
    body=body_of(c);body['students'].append({'student_id':student(t,campus),'level':'B','note':'','weekly_slots':[{'weekday':1,'slot_number':3}]})
    assert t.patch(f"/api/teacher/classes/{c['id']}",json=body).status_code==200
    historic=workbook(a,f'view=occupancy&broadcast_id={bid}');current=workbook(a,'view=occupancy')
    assert historic['教师满班率']['G2'].value==1 and current['教师满班率']['G2'].value==2
    assert workbook(a,f'view=students&broadcast_id={bid}')['学生列表'].max_row==2
