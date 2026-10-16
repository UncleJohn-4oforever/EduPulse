from io import BytesIO
from pathlib import Path
import os
from datetime import datetime
from zoneinfo import ZoneInfo
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from PIL import Image, ImageDraw, ImageFont

def local(v):
    return datetime.fromisoformat(v).astimezone(ZoneInfo('Asia/Shanghai')).strftime('%Y-%m-%d %H:%M') if v else '—'
def safe(v):
    # Prevent formula injection from user-controlled names/notes.
    return "'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v
def slot_text(slots):
    return '；'.join(f"周{'一二三四五六日'[s['weekday']-1]} {s['slot_label']}"+(f" {s['start_time']}-{s['end_time']}" if s['start_time'] else '') for s in slots)
def workload_note(data):
    return '周负荷按学生档位在同班、同星期、同档位去重' if data.get('workload_basis')=='student-slot-union' else '周负荷按旧版班级固定时段数'
def legacy_xlsx(data):
    wb=Workbook(); wb.remove(wb.active)
    def sheet(name,headers,rows,percentage_columns=()):
        ws=wb.create_sheet(name); ws.append(headers)
        for row in rows: ws.append([safe(v) for v in row])
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]: cell.font=Font(color='FFFFFF',bold=True); cell.fill=PatternFill('solid',fgColor='0F766E')
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width=min(50,max(14,max(len(str(c.value or '')) for c in col)+2))
            for cell in col: cell.alignment=Alignment(vertical='top',wrap_text=True)
        for r in range(2,ws.max_row+1):
            for c in percentage_columns:
                if isinstance(ws.cell(r,c).value,(int,float)): ws.cell(r,c).number_format='0.00%'
    rows=[]
    for t in data['teachers']:
        row=[t['display_name'],t['username'],t['campus'],t['groups']['1']['classes']]
        for cap in ('3','6'):
            g=t['groups'][cap]; row.extend([g['classes'],g['seats'],g['capacity'],g['rate'] if g['rate'] is not None else 'N/A',g['vacancies']])
        row.extend([t['small_rate'] if t['small_rate'] is not None else 'N/A',t.get('confirmation','当前数据')]); rows.append(row)
    sheet('教师满班率',['教师','账号','校区','一对一班数','一对三班数','一对三席位','一对三容量','一对三满班率','一对三空位','一对六班数','一对六席位','一对六容量','一对六满班率','一对六空位','小班综合满班率','本轮确认'],rows,(8,13,15))
    sheet('教师周负荷',['教师','校区','一对一周课次','一对三周课次','一对六周课次','总周课次','班级数'],[[t['display_name'],t['campus'],*[t['groups'][str(c)]['weekly_lessons'] for c in (1,3,6)],t['weekly_lessons'],t['class_count']] for t in data['teachers']])
    sheet('学生科目教师评价明细',['学生编号','姓名','年级','去重科目数','科目','任课教师','班级','校区','程度','文字评价','评价更新时间','该生上课档位'],[[s['id'],s['display_name'],s['grade'],s['subject_count'],d['subject'],d['teacher'],d['class_name'],d['campus'],d['level'],d['note'],local(d['updated_at']),slot_text(d.get('weekly_slots',[]))] for s in data['students'] for d in s['details']])
    sheet('班级明细',['班级编号','教师','名称','科目','年级','容量','学生席位','每周课次','时段'],[[c['id'],t['display_name'],c['display_name'],c['subject'],c['grade'],c['capacity'],len(c['students']),len(c['weekly_slots']),'；'.join(f"周{'一二三四五六日'[s['weekday']-1]} {s['slot_label']} {s['start_time'] or ''}-{s['end_time'] or ''}" for s in c['weekly_slots'])] for t in data['teachers'] for c in t['classes']])
    progress=data.get('broadcast',{}).get('targets',[])
    sheet('填报完成清单',['教师','账号','确认状态','确认时间'],[[t['display_name'],t['username'],'已确认' if t['confirmed_at'] else '未经本轮确认',local(t['confirmed_at'])] for t in progress])
    sheet('统计说明',['项目','内容'],[['生成时间（北京时间）',local(data['generated_at'])],['范围',str(data.get('scope', '全部校区'))],['广播',data.get('broadcast',{}).get('title','当前实时数据')],['口径','小班满班率 = 已用席位 / 一对三与一对六总容量；'+workload_note(data)+'；学生科目去重。']])
    out=BytesIO(); wb.save(out); return out.getvalue()
def font_path():
    candidates=[os.environ.get('EDUPULSE_FONT',''),str(Path(__file__).parents[2]/'assets/NotoSansSC.ttf'),'C:/Windows/Fonts/msyh.ttc','/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc']
    return next((p for p in candidates if p and Path(p).is_file()),None)
def legacy_png(data):
    path=font_path()
    if not path: raise RuntimeError('未找到中文字体，请配置 EDUPULSE_FONT')
    font=ImageFont.truetype(path,25); small=ImageFont.truetype(path,20); title=ImageFont.truetype(path,42)
    rows=data['teachers']; height=330+max(1,len(rows))*65
    image=Image.new('RGB',(1500,height),'#f3f7f6'); draw=ImageDraw.Draw(image)
    draw.rectangle((0,0,1500,128),fill='#103f3a'); draw.text((48,28),'EduPulse · 教师教学数据看板',font=title,fill='white')
    draw.text((48,90),f"统计时间（北京时间）：{local(data['generated_at'])}",font=small,fill='#b9ded7')
    scope=data.get('scope',{}); draw.text((48,145),f"校区范围：{scope.get('campus_label','全部校区')}   科目：{scope.get('subject_label','全部')}   搜索：{scope.get('q','') or '全部'}",font=small,fill='#334155')
    b=data.get('broadcast'); label=f"广播：{b['title']} · 已确认 {sum(bool(t['confirmed_at']) for t in b['targets'])}/{len(b['targets'])}" if b else '当前实时数据 · 未关联广播快照'
    draw.text((48,180),label[:95],font=small,fill='#334155')
    xs=[48,280,440,660,880,1110,1300]; headers=['教师 / 校区','1对1 班/周次','1对3 席位/容量','1对6 席位/容量','小班综合','总周课次','本轮核对']
    draw.rectangle((35,225,1465,273),fill='#dcece8')
    for x,h in zip(xs,headers): draw.text((x,235),h,font=small,fill='#103f3a')
    rate=lambda r: f'{r:.2%}' if r is not None else '—'
    for i,t in enumerate(rows):
        y=282+i*65
        draw.rectangle((35,y-2,1465,y+57),fill='white' if i%2==0 else '#eaf2ef')
        a=t['groups']['1']; g3=t['groups']['3']; g6=t['groups']['6']
        values=[f"{t['display_name'][:9]} / {t['campus'][:6]}",f"{a['classes']} / {a['weekly_lessons']}",f"{g3['seats']}/{g3['capacity']}  {rate(g3['rate'])}",f"{g6['seats']}/{g6['capacity']}  {rate(g6['rate'])}",rate(t['small_rate']),str(t['weekly_lessons']),t.get('confirmation','当前数据')]
        for x,v in zip(xs,values): draw.text((x,y+12),v,font=small,fill='#243e39')
    if not rows: draw.text((48,290),'暂无教师数据',font=font,fill='#334155')
    draw.text((48,height-40),'口径：席位按班级关联计数；小班按容量加权（不含一对一）；'+workload_note(data)+'。',font=small,fill='#475569')
    out=BytesIO(); image.save(out,format='PNG'); return out.getvalue()


VIEW_NAMES={'occupancy':'教师满班率','workload':'教师负荷','students':'学生列表'}

def average(seats,classes):
    return seats/classes if classes else None

def page_tables(data,view):
    """Return only the current report's tables, with all filtered records."""
    if view=='students':
        overview=[[s['id'],s['display_name'],s['grade'],s['subject_count'],' / '.join(s['subjects']),len(s['details'])] for s in data['students']]
        details=[[s['id'],s['display_name'],s['grade'],s['subject_count'],d['subject'],d['teacher'],d['class_name'],d['campus'],d['level'],d['note'],local(d['updated_at']),slot_text(d.get('weekly_slots',[]))] for s in data['students'] for d in s['details']]
        return [('学生列表',['学生编号','姓名','年级','去重科目数','科目','关联班级数'],overview,(),()),('学生评价与档位明细',['学生编号','姓名','年级','去重科目数','科目','任课教师','班级','校区','程度','文字评价','评价更新时间','该生上课档位'],details,(),())]
    if view=='workload':
        rows=[[t['display_name'],t['username'],t['campus'],*[t['groups'][str(c)]['weekly_lessons'] for c in (1,3,6)],t['weekly_lessons'],15,f"{t['weekly_lessons']}/15",t['weekly_lessons']/15,t['class_count'],t.get('confirmation','当前数据')] for t in data['teachers']]
        return [('教师负荷',['教师','账号','校区','一对一周课次','一对三周课次','一对六周课次','总周课次','基准课次','本周课次/基准','负荷率','班级数','本轮核对'],rows,(10,),())]
    rows=[]
    for t in data['teachers']:
        row=[t['display_name'],t['username'],t['campus'],t['groups']['1']['classes'],t['groups']['1']['weekly_lessons']]
        for cap in ('3','6'):
            g=t['groups'][cap];row.extend([g['classes'],g['seats'],g['capacity'],g['rate'],average(g['seats'],g['classes']),g['vacancies']])
        row.extend([t['small_rate'],average(t['small_seats'],t['groups']['3']['classes']+t['groups']['6']['classes']),t['small_capacity']-t['small_seats'],t.get('confirmation','当前数据')]);rows.append(row)
    headers=['教师','账号','校区','一对一班数','一对一周课次']
    for cap in ('三','六'): headers.extend([f'一对{cap}班数',f'一对{cap}席位',f'一对{cap}容量',f'一对{cap}满班率',f'一对{cap}班容（人/班）',f'一对{cap}空位'])
    headers.extend(['小班综合满班率','小班综合班容（人/班）','小班空位','本轮核对'])
    summary=[]
    for label,caps in [('一对三',('3',)),('一对六',('6',)),('综合平均',('3','6'))]:
        classes=sum(t['groups'][c]['classes'] for t in data['teachers'] for c in caps)
        seats=sum(t['groups'][c]['seats'] for t in data['teachers'] for c in caps)
        capacity=sum(t['groups'][c]['capacity'] for t in data['teachers'] for c in caps)
        summary.append([label,classes,seats,capacity,seats/capacity if capacity else None,average(seats,classes)])
    return [('教师满班率',headers,rows,(9,15,18),(10,16,19)),('满班率汇总',['班型','班级数','席位','总容量','满班率','班容（人/班）'],summary,(5,),(6,))]


def xlsx(data,view='all'):
    if view=='all': return legacy_xlsx(data)
    wb=Workbook();wb.remove(wb.active)
    tables=page_tables(data,view)
    scope=data.get('scope',{})
    tables.append(('统计说明',['项目','内容'],[
        ['导出页面',VIEW_NAMES[view]],['导出范围','当前筛选下的全部记录（包含翻页后的记录）'],
        ['生成时间（北京时间）',local(data['generated_at'])],['校区',scope.get('campus_label','全部校区')],['科目',scope.get('subject_label','全部科目')],
        ['筛选条件',str(scope)],['数据来源',data.get('broadcast',{}).get('title','当前实时数据')],
        ['口径','小班满班率按容量加权，不含一对一；班容=席位÷班级数；负荷率=本周课次÷15，可超过100%；'+workload_note(data)+'；学生科目去重。']],(),()))
    for name,headers,rows,percentages,decimals in tables:
        ws=wb.create_sheet(name);ws.append(headers)
        for row in rows:ws.append([safe(v) if v is not None else '—' for v in row])
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.font=Font(color='FFFFFF',bold=True);c.fill=PatternFill('solid',fgColor='0F766E')
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width=min(55,max(16,max(len(str(c.value or '')) for c in col)+2))
            for cell in col:cell.alignment=Alignment(vertical='top',wrap_text=True)
        for r in range(2,ws.max_row+1):
            for c in (*percentages,*decimals):
                if isinstance(ws.cell(r,c).value,(float,int)):ws.cell(r,c).number_format='0.00%' if c in percentages else '0.00'
    out=BytesIO();wb.save(out);return out.getvalue()


def png(data,view='all'):
    if view=='all':return legacy_png(data)
    path=font_path()
    if not path:raise RuntimeError('未找到中文字体，请配置 EDUPULSE_FONT')
    small=ImageFont.truetype(path,21);title=ImageFont.truetype(path,38)
    rate=lambda value:f'{value:.2%}' if value is not None else '—'
    size=lambda seats,classes:f'{seats/classes:.2f}' if classes else '—'
    if view=='students':
        widths=[110,190,180,120,360,440];headers=['编号','姓名','年级','科目数','科目','任课教师 / 班级']
        rows=[[str(s['id']),s['display_name'],s['grade'],str(s['subject_count']),' / '.join(s['subjects']),'；'.join(f"{d['teacher']} · {d['class_name']}" for d in s['details']) or '尚未关联班级'] for s in data['students']]
    elif view=='workload':
        widths=[260,170,170,170,190,160,120,160];headers=['教师 / 校区','一对一周课次','一对三周课次','一对六周课次','本周课次 / 15','负荷率','班级数','本轮核对']
        rows=[[t['display_name']+' / '+t['campus'],*[str(t['groups'][str(c)]['weekly_lessons']) for c in (1,3,6)],f"{t['weekly_lessons']}/15",rate(t['weekly_lessons']/15),str(t['class_count']),t.get('confirmation','当前数据')] for t in data['teachers']]
    else:
        widths=[260,180,300,300,300,130,160];headers=['教师 / 校区','一对一 班 / 周次','一对三 满班率 · 班容','一对六 满班率 · 班容','综合满班率 · 班容','小班空位','本轮核对']
        rows=[]
        for t in data['teachers']:
            a=t['groups']['1'];groups=[t['groups'][c] for c in ('3','6')]
            rows.append([t['display_name']+' / '+t['campus'],f"{a['classes']} / {a['weekly_lessons']}",*[rate(g['rate'])+' · '+size(g['seats'],g['classes'])+'人/班' for g in groups],rate(t['small_rate'])+' · '+size(t['small_seats'],sum(g['classes'] for g in groups))+'人/班',str(t['small_capacity']-t['small_seats']),t.get('confirmation','当前数据')])
    # Wrap long names and text rather than allowing columns to overlap.
    def wrap(value,width):
        lines=['']
        for ch in str(value):
            if small.getlength(lines[-1]+ch)>width and lines[-1]:lines.append(ch)
            else:lines[-1]+=ch
        return lines
    wrapped=[[[*wrap(v,w-20)] for v,w in zip(row,widths)] for row in rows]
    heights=[max(len(cell) for cell in row)*30+22 for row in wrapped]
    width=sum(widths)+70;height=330+sum(heights or [60])
    image=Image.new('RGB',(width,height),'#f3f7f6');draw=ImageDraw.Draw(image)
    draw.rectangle((0,0,width,120),fill='#103f3a');draw.text((35,22),'EduPulse · '+VIEW_NAMES[view],font=title,fill='white')
    draw.text((35,78),'统计时间（北京时间）：'+local(data['generated_at']),font=small,fill='#b9ded7')
    scope=data.get('scope',{});label=f"{scope.get('campus_label','全部校区')} · {scope.get('subject_label','全部科目')} · 共{len(rows)}条 · 筛选后全部记录"
    draw.text((35,140),label,font=small,fill='#334155')
    draw.text((35,178),data.get('broadcast',{}).get('title','当前实时数据')[:70],font=small,fill='#334155')
    draw.rectangle((35,220,width-35,268),fill='#dcece8');x=45
    for label,w in zip(headers,widths):draw.text((x,230),label,font=small,fill='#103f3a');x+=w
    y=278
    for i,(row,h) in enumerate(zip(wrapped,heights)):
        draw.rectangle((35,y,width-35,y+h),fill='white' if i%2==0 else '#eaf2ef');x=45
        for cell,w in zip(row,widths):draw.multiline_text((x,y+10),'\n'.join(cell),font=small,spacing=6,fill='#243e39');x+=w
        y+=h
    if not rows:draw.text((45,y+10),'暂无符合条件的数据',font=small,fill='#334155')
    draw.text((35,height-38),'班容=席位÷班级数；小班按容量加权；负荷=本周课次÷15，可超过100%；学生科目去重。',font=small,fill='#475569')
    out=BytesIO();image.save(out,format='PNG');return out.getvalue()
